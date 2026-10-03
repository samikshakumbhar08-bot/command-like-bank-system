from openpyxl import Workbook, load_workbook
import os

FILE = "bank.xlsx"

# Create Excel file if it does not exist
if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws.append(["Account No", "Name", "Mobile", "Balance"])
    wb.save(FILE)


def create_account():
    wb = load_workbook(FILE)
    ws = wb["Accounts"]

    acc_no = input("Enter Account Number: ")

    # Check account already exists
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == acc_no:
            print("Account already exists!")
            wb.close()
            return

    name = input("Enter Account Holder Name: ")
    mobile = input("Enter Mobile Number: ")
    balance = float(input("Enter Initial Deposit: "))

    ws.append([acc_no, name, mobile, balance])
    wb.save(FILE)
    wb.close()

    print("Account Created Successfully!")


def deposit_money():
    wb = load_workbook(FILE)
    ws = wb["Accounts"]

    acc_no = input("Enter Account Number: ")
    amount = float(input("Enter Amount to Deposit: "))

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == acc_no:
            row[3].value += amount
            wb.save(FILE)
            wb.close()

            print("Amount Deposited Successfully!")
            print("Current Balance:", row[3].value)
            return

    print("Account not found!")
    wb.close()


def withdraw_money():
    wb = load_workbook(FILE)
    ws = wb["Accounts"]

    acc_no = input("Enter Account Number: ")
    amount = float(input("Enter Amount to Withdraw: "))

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == acc_no:

            if amount > row[3].value:
                print("Insufficient Balance!")
            else:
                row[3].value -= amount
                wb.save(FILE)

                print("Amount Withdrawn Successfully!")
                print("Current Balance:", row[3].value)

            wb.close()
            return

    print("Account not found!")
    wb.close()


def show_balance():
    wb = load_workbook(FILE)
    ws = wb["Accounts"]

    acc_no = input("Enter Account Number: ")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == acc_no:
            print("Account Number:", row[0])
            print("Current Balance:", row[3])
            wb.close()
            return

    print("Account not found!")
    wb.close()


def account_details():
    wb = load_workbook(FILE)
    ws = wb["Accounts"]

    acc_no = input("Enter Account Number: ")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == acc_no:
            print("\n----- ACCOUNT DETAILS -----")
            print("Account Number:", row[0])
            print("Account Holder:", row[1])
            print("Mobile Number:", row[2])
            print("Balance:", row[3])

            wb.close()
            return

    print("Account not found!")
    wb.close()


# Main Menu
while True:
    print("\n==============================")
    print("     BANK MANAGEMENT SYSTEM")
    print("==============================")
    print("1. Create Account")
    print("2. Deposit Money")
    print("3. Withdraw Money")
    print("4. Show Balance")
    print("5. Account Details")
    print("6. Exit")
    print("==============================")

    choice = input("Enter your choice: ")

    if choice == "1":
        create_account()

    elif choice == "2":
        deposit_money()

    elif choice == "3":
        withdraw_money()

    elif choice == "4":
        show_balance()

    elif choice == "5":
        account_details()

    elif choice == "6":
        print("Thank you for using Bank Management System!")
        break

    else:
        print("Invalid Choice!")
